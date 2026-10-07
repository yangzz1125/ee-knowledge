"""以 9.22 汇报的 PPT 为模板，生成第二次汇报（对象与数据结构设计）的 PPT。

用法：
    python build_deck.py <模板.pptx> <公式图片.png> <pptx skill 的 scripts 目录> <输出.pptx>

封面、目录、过渡页和内容页的页眉沿用模板；内容页正文全部用原生形状和文本框绘制，
只有公式是图片（formula.tex 编译后用 pdftoppm -r 600 导出）。
正文坐标沿用原设计稿的画布：宽 31cm、高 14.3cm、原点在左下角，对应模板中正文图片的区域。
"""
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

TEMPLATE, FORMULA_PNG, SKILL_SCRIPTS, OUTPUT = (Path(a) for a in sys.argv[1:5])

TITLE = "电磁场与波AI知识图谱对象与数据结构设计"
DATE = "2026年10月7日"
SECTIONS = ["本次工作内容", "知识对象设计", "知识点内容设计", "知识关系设计", "数据校验", "后续计划"]
NUMERALS = ["一", "二", "三", "四", "五", "六"]
# (所属章节序号, 小标题)，顺序与正文绘制函数 BODIES 对应
CONTENT = [
    (0, "分阶段建设路线与本次任务"),
    (1, "对象的划分"),
    (1, "章节与公式的定位"),
    (2, "共有内容与各类专有内容"),
    (2, "公式与变量"),
    (3, "九种关系类型"),
    (3, "跨章节的知识联系示例"),
    (3, "整理关系时的约定"),
    (4, "检查流程与规则"),
    (5, "下一阶段工作"),
]


# ============================================================ 模板处理（原始 XML）

def run(*args):
    subprocess.run([sys.executable, *map(str, args)], check=True, cwd=SKILL_SCRIPTS)


def read(p: Path) -> str:
    return p.read_text(encoding="utf-8")


def write(p: Path, s: str):
    p.write_text(s, encoding="utf-8")


def sub_once(pattern: str, repl: str, s: str, flags=0) -> str:
    out, n = re.subn(pattern, repl, s, count=1, flags=flags)
    if n != 1:
        raise ValueError(f"模板中找不到：{pattern}")
    return out


def esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def set_shape_text(xml: str, shape_name: str, text: str) -> str:
    """把指定文本框中的所有文字替换为 text：保留第一个 run 的格式，删掉其余 run。"""
    m = re.search(r'<p:sp>(?:(?!</p:sp>).)*?name="%s".*?</p:sp>' % re.escape(shape_name), xml, re.S)
    if not m:
        raise ValueError(f"找不到形状：{shape_name}")
    sp = m.group(0)
    runs = re.findall(r"<a:r>.*?</a:r>", sp, re.S)
    first = re.sub(r"<a:t>.*?</a:t>", f"<a:t>{esc(text)}</a:t>", runs[0], flags=re.S)
    new_sp = sp.replace(runs[0], first, 1)
    for r in runs[1:]:
        new_sp = new_sp.replace(r, "", 1)
    return xml.replace(sp, new_sp, 1)


def build_from_template(unpacked: Path) -> list[int]:
    """完成模板部分，返回各内容页在演示文稿中的序号（从 0 起）。"""
    slides = unpacked / "ppt" / "slides"

    # 1. 结构：复制过渡页（slide3）和内容页（slide4），再按目标顺序排列
    def newest():
        return sorted(slides.glob("slide*.xml"), key=lambda p: int(re.findall(r"\d+", p.name)[0]))[-1].name

    dividers = ["slide3.xml"]
    contents = ["slide4.xml"]
    for _ in range(len(SECTIONS) - 1):
        run("add_slide.py", unpacked, "slide3.xml")
        dividers.append(newest())
    for _ in range(len(CONTENT) - 1):
        run("add_slide.py", unpacked, "slide4.xml")
        contents.append(newest())

    order = ["slide1.xml", "slide2.xml"]
    ci = 0
    for si in range(len(SECTIONS)):
        order.append(dividers[si])
        while ci < len(CONTENT) and CONTENT[ci][0] == si:
            order.append(contents[ci])
            ci += 1

    pres_rels = read(unpacked / "ppt" / "_rels" / "presentation.xml.rels")
    rid_of = {t: rid for rid, t in re.findall(r'Id="(rId\d+)"[^>]*Target="slides/(slide\d+\.xml)"', pres_rels)}
    rid_of.update({t: rid for t, rid in re.findall(r'Target="slides/(slide\d+\.xml)"[^>]*Id="(rId\d+)"', pres_rels)})
    pres_path = unpacked / "ppt" / "presentation.xml"
    pres = read(pres_path)
    ids = {rid: sid for sid, rid in re.findall(r'<p:sldId id="(\d+)" r:id="(rId\d+)"/>', pres)}
    new_lst = "".join(f'<p:sldId id="{ids[rid_of[s]]}" r:id="{rid_of[s]}"/>' for s in order)
    pres = sub_once(r"<p:sldIdLst>.*?</p:sldIdLst>", f"<p:sldIdLst>{new_lst}</p:sldIdLst>", pres, re.S)
    write(pres_path, pres)
    run("clean.py", unpacked)

    # 2. 封面
    p = slides / "slide1.xml"
    x = set_shape_text(read(p), "文本框 10", TITLE)
    x = x.replace('sz="5000"', 'sz="4000"')
    x = set_shape_text(x, "文本框 2", DATE)
    write(p, x)

    # 3. 目录：原有 4 个条目按 y 排序后复用，再补 2 个
    p = slides / "slide2.xml"
    x = read(p)
    x = re.sub(r'<p:sp>(?:(?!</p:sp>).)*?name="文本框 10".*?</p:sp>', "", x, count=1, flags=re.S)
    toc = []
    for m in re.finditer(r"<p:sp>(?:(?!</p:sp>).)*?</p:sp>", x, re.S):
        sp = m.group(0)
        if re.search(r"<a:t>[一二三四]、", sp):
            toc.append((int(re.search(r'<a:off x="\d+" y="(\d+)"', sp).group(1)), sp))
    toc.sort()
    proto = toc[0][1]
    for _, sp in toc:
        x = x.replace(sp, "", 1)
    max_id = max(int(i) for i in re.findall(r'<p:cNvPr id="(\d+)"', x))
    items = []
    for k, sec in enumerate(SECTIONS):
        sp = re.sub(r'<p:cNvPr id="\d+" name="[^"]*"', f'<p:cNvPr id="{max_id + 1 + k}" name="目录 {k + 1}"', proto, count=1)
        sp = re.sub(r'<a:off x="(\d+)" y="\d+"/>', rf'<a:off x="3133193" y="{1600000 + k * 750000}"/>', sp, count=1)
        runs = re.findall(r"<a:r>.*?</a:r>", sp, re.S)
        first = re.sub(r"<a:t>.*?</a:t>", f"<a:t>{NUMERALS[k]}、{sec}</a:t>", runs[0], flags=re.S)
        sp = sp.replace(runs[0], first, 1)
        for r in runs[1:]:
            sp = sp.replace(r, "", 1)
        items.append(sp)
    x = x.replace("</p:spTree>", "".join(items) + "</p:spTree>", 1)
    write(p, x)

    # 4. 章节过渡页
    for k, name in enumerate(dividers):
        p = slides / name
        write(p, set_shape_text(read(p), "文本框 12", f"{NUMERALS[k]}、{SECTIONS[k]}"))

    # 5. 内容页：填章节名和小标题，删掉模板的正文图片（正文之后用原生形状绘制）
    for name, (sec, title) in zip(contents, CONTENT):
        p = slides / name
        x = set_shape_text(read(p), "文本框 20", SECTIONS[sec])
        x = set_shape_text(x, "文本框 7", title)
        x = sub_once(r'<p:pic>(?:(?!</p:pic>).)*?name="图片 17".*?</p:pic>', "", x, re.S)
        write(p, x)
        rp = slides / "_rels" / f"{name}.rels"
        write(rp, sub_once(r'<Relationship [^>]*Id="rId4"[^>]*/>', "", read(rp)))

    # 6. 清空保留页上的旧演讲者备注
    for notes in (unpacked / "ppt" / "notesSlides").glob("notesSlide*.xml"):
        write(notes, re.sub(r"<a:t>[^<]*</a:t>", "<a:t></a:t>", read(notes)))

    run("clean.py", unpacked)
    return [order.index(c) for c in contents]


# ============================================================ 正文绘制（python-pptx）

BLUE, TEAL, PURPLE, ORANGE, RED, GREEN = "1F6FEB", "0EA5A0", "7C3AED", "F59E0B", "EF4444", "10B981"
INK, TEXT, GRAY, LINE, NAVY, WHITE = "1E293B", "334155", "64748B", "D7E0EC", "0259A0", "FFFFFF"
FONT = "微软雅黑"

# 画布（cm，y 轴向上）到幻灯片（EMU）的映射：画布左上角对应模板正文图片区域的左上角
X0, Y0, H = 467334, 1391049, 14.3
CM = 360000


def ex(x):
    return Emu(round(X0 + x * CM))


def ey(y):
    return Emu(round(Y0 + (H - y) * CM))


def el(d):
    return Emu(round(d * CM))


def mix(c, p, base=WHITE):
    """c 占 p、base 占 1-p 的混色，对应 TikZ 的 c!p!base。"""
    a, b = bytes.fromhex(c), bytes.fromhex(base)
    return "".join(f"{round(x * p + y * (1 - p)):02X}" for x, y in zip(a, b))


def pt2cm(v):
    return v * 0.03528


def tw(s, size):
    """估算单行文字宽度（cm）：汉字按 1em，空格 0.3em，其他字符 0.55em。"""
    em = sum(1.0 if ord(ch) > 0x2E80 else 0.3 if ch == " " else 0.55 for ch in s)
    return pt2cm(em * size)


class Box:
    """画布上的矩形：l 为左边，t 为上边（y 轴向上）。"""

    def __init__(self, l, t, w, h):
        self.l, self.t, self.w, self.h = l, t, w, h

    cx = property(lambda s: s.l + s.w / 2)
    cy = property(lambda s: s.t - s.h / 2)
    r = property(lambda s: s.l + s.w)
    b = property(lambda s: s.t - s.h)
    west = property(lambda s: (s.l, s.cy))
    east = property(lambda s: (s.r, s.cy))
    north = property(lambda s: (s.cx, s.t))
    south = property(lambda s: (s.cx, s.b))


def place(anchor, x, y, w, h):
    """按 TikZ 的锚点（c / n / s / w / e / nw / ne / sw / se）放置 w×h 的矩形。"""
    dx = {"w": 0, "e": -w}.get(anchor[-1], -w / 2)
    dy = {"n": 0, "s": h}.get(anchor[0], h / 2)
    return Box(x + dx, y + dy, w, h)


class R:
    """一段文字及其格式；未给出的格式沿用段落或文本框的默认值。"""

    def __init__(self, text, color=None, bold=None, size=None, italic=False, latin=None):
        self.text, self.color, self.bold, self.size, self.italic, self.latin = text, color, bold, size, italic, latin


class P:
    """一个段落：若干 R 或字符串；bullet 为项目符号颜色。"""

    def __init__(self, *runs, bullet=None, before=0):
        self.runs = [r if isinstance(r, R) else R(r) for r in runs]
        self.bullet, self.before = bullet, before


SHADOW = (f'<a:effectLst {nsdecls("a")}><a:outerShdw blurRad="50800" dist="19050" dir="5400000" '
          'algn="t" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="10000"/></a:srgbClr>'
          '</a:outerShdw></a:effectLst>')

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
VANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def fill_text(tf, paras, size=14, color=TEXT, bold=False, align="l", line=None, anchor="t", wrap=True,
              margin=(0, 0, 0, 0)):
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = VANCHOR[anchor]
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = (el(m) for m in margin)
    if isinstance(paras, (str, P)):
        paras = [paras]
    for i, para in enumerate(paras):
        para = para if isinstance(para, P) else P(para)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if line:
            p.line_spacing = Pt(line)
        if para.before:
            p.space_before = Pt(para.before)
        if para.bullet:
            pPr = p._p.get_or_add_pPr()
            ind = el(pt2cm(size) * 1.25)
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            pPr.append(parse_xml(f'<a:buClr {nsdecls("a")}><a:srgbClr val="{para.bullet}"/></a:buClr>'))
            pPr.append(parse_xml(f'<a:buSzPct {nsdecls("a")} val="80000"/>'))
            pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="{FONT}"/>'))
            pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="●"/>'))
        for rs in para.runs:
            r = p.add_run()
            r.text = rs.text
            f = r.font
            f.size = Pt(rs.size or size)
            f.bold = bold if rs.bold is None else rs.bold
            if rs.italic:
                f.italic = True
            f.color.rgb = RGBColor.from_string(rs.color or color)
            rPr = r._r.get_or_add_rPr()
            rPr.set("lang", "zh-CN")
            rPr.set("altLang", "en-US")
            etree.SubElement(rPr, qn("a:latin")).set("typeface", rs.latin or FONT)
            etree.SubElement(rPr, qn("a:ea")).set("typeface", FONT)


class Canvas:
    def __init__(self, slide):
        self.shapes = slide.shapes
        self.n = 0

    def _finish(self, shp, kind):
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        self.n += 1
        shp.name = f"{kind} {self.n}"
        return shp

    # ---- 基本图形
    def shape(self, kind, prst, box, fill=None, line=None, lw=0.8, radius=None, shadow=False):
        shp = self._finish(self.shapes.add_shape(prst, ex(box.l), ey(box.t), el(box.w), el(box.h)), kind)
        if radius is not None:
            shp.adjustments[0] = min(0.5, radius / min(box.w, box.h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = RGBColor.from_string(line)
            shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        if shadow:
            shp._element.spPr.append(parse_xml(SHADOW))
        return shp

    def card(self, box):
        return self.shape("卡片", MSO_SHAPE.ROUNDED_RECTANGLE, box, WHITE, LINE, 0.8, radius=0.247, shadow=True)

    def tint(self, box, c, p=0.07):
        return self.shape("底色", MSO_SHAPE.ROUNDED_RECTANGLE, box, mix(c, p), radius=0.247)

    def text(self, box, paras, kind="文字", fill=None, **kw):
        tb = self._finish(self.shapes.add_textbox(ex(box.l), ey(box.t), el(box.w), el(box.h)), kind)
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = RGBColor.from_string(fill)
        fill_text(tb.text_frame, paras, **kw)
        return tb

    def at(self, anchor, x, y, paras, size=14, w=None, h=None, line=None, **kw):
        """按锚点放一段文字；不给 w 时按文字估算宽度且不换行。"""
        plist = [paras] if isinstance(paras, (str, P)) else paras
        plain = ["".join(r.text for r in p.runs) if isinstance(p, P) else p for p in plist]
        wrap = w is not None
        w = w if wrap else max(tw(s, size) for s in plain) + 0.3
        h = h or pt2cm(line or size * 1.3) * len(plist)
        horiz = {"w": "l", "e": "r"}.get(anchor[-1], "c")
        vert = {"n": "t", "s": "b"}.get(anchor[0], "m")
        return self.text(place(anchor, x, y, w, h), paras, size=size, line=line, wrap=wrap,
                         align=kw.pop("align", horiz), anchor=vert, **kw)

    def num(self, x, y, label, c, d=0.95, size=15):
        shp = self.shape("序号", MSO_SHAPE.OVAL, place("c", x, y, d, d), c)
        fill_text(shp.text_frame, label, size=size, color=WHITE, bold=True, align="c", anchor="m", wrap=False)
        return shp

    def chip(self, anchor, x, y, label, c, size=12.5, fill=None, color=None):
        box = place(anchor, x, y, tw(label, size) + 0.5, pt2cm(size * 1.05 + 8))
        shp = self.shape("标签", MSO_SHAPE.ROUNDED_RECTANGLE, box, fill or mix(c, 0.12), radius=0.141)
        fill_text(shp.text_frame, label, size=size, color=color or mix(c, 0.75, "000000"), bold=True,
                  align="c", anchor="m", wrap=False)
        return box

    def ent(self, x, y, paras, c, w=None, h=1.15, size=14):
        """知识点节点：带边框的浅色圆角矩形，文字居中。"""
        plist = [paras] if isinstance(paras, (str, P)) else paras
        first = paras if isinstance(paras, str) else "".join(r.text for r in plist[0].runs) if isinstance(plist[0], P) else plist[0]
        box = place("c", x, y, w or max(3.4, tw(first, size) + 0.6), h)
        shp = self.shape("知识点", MSO_SHAPE.ROUNDED_RECTANGLE, box, mix(c, 0.10), c, 1.0, radius=0.176)
        fill_text(shp.text_frame, paras, size=size, color=TEXT, align="c", anchor="m", wrap=False)
        return box

    # ---- 线条
    def _ln(self, shp, color, lw, arrow, dash, cap):
        shp.line.color.rgb = RGBColor.from_string(color)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
        ln = shp._element.spPr.find(qn("a:ln"))
        if cap:
            ln.set("cap", "rnd")
        if arrow:
            ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="stealth" w="lg" len="lg"/>'))

    def line(self, p1, p2, color=GRAY, lw=1.1, arrow=False, dash=False, cap=False):
        shp = self._finish(self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ex(p1[0]), ey(p1[1]),
                                                     ex(p2[0]), ey(p2[1])), "连线")
        self._ln(shp, color, lw, arrow, dash, cap)
        return shp

    def poly(self, pts, color=GRAY, lw=1.1, arrow=True, dash=False):
        fb = self.shapes.build_freeform(ex(pts[0][0]), ey(pts[0][1]), scale=1.0)
        fb.add_line_segments([(ex(x), ey(y)) for x, y in pts[1:]], close=False)
        shp = self._finish(fb.convert_to_shape(), "折线")
        shp.fill.background()
        self._ln(shp, color, lw, arrow, dash, False)
        return shp

    def label(self, x, y, s, size=12):
        """连线上的说明文字（白底，盖住线条）。"""
        box = place("c", x, y, tw(s, size) + 0.25, pt2cm(size * 1.35))
        return self.text(box, s, kind="连线说明", fill=WHITE, size=size, color=INK, align="c", anchor="m", wrap=False)

    def link(self, a, b, text=None, color=GRAY, lw=1.1, arrow=True, dash=False, shift=(0, 0)):
        """连接两个矩形的中心，端点裁到矩形边上；text 标在线段中点。"""
        def clip(box, ox, oy):
            dx, dy = ox - box.cx, oy - box.cy
            t = min(box.w / 2 / abs(dx) if dx else 1e9, box.h / 2 / abs(dy) if dy else 1e9)
            return box.cx + dx * t, box.cy + dy * t
        p1, p2 = clip(a, b.cx, b.cy), clip(b, a.cx, a.cy)
        self.line(p1, p2, color, lw, arrow, dash)
        if text:
            self.label((p1[0] + p2[0]) / 2 + shift[0], (p1[1] + p2[1]) / 2 + shift[1], text)

    def bottombar(self, y, label, text):
        bar = Box(0.4, y + 1.25, 30.2, 1.25)
        self.tint(bar, NAVY)
        self.at("w", bar.l + 0.5, bar.cy, label, size=17, color=NAVY, bold=True)
        self.line((bar.l + 3.6, bar.cy - 0.35), (bar.l + 3.6, bar.cy + 0.35), mix(NAVY, 0.4), 1.0)
        self.at("w", bar.l + 4.0, bar.cy, text, size=14)

    def heading(self, card, n, label, c, dy=0.9, title_size=18):
        """卡片左上角的序号圆点和标题。"""
        y = card.t - dy
        self.num(card.l + 0.9, y, n, c)
        self.at("w", card.l + 1.7, y, label, size=title_size, color=INK, bold=True)


# ---- 各页正文

def output_box(c, card, col, text):
    """卡片底部的「产出」框。"""
    box = Box(card.l + 0.3, card.b + 0.3 + 1.75, card.w - 0.6, 1.75)
    c.tint(box, col, 0.10)
    c.text(Box(box.l + 0.3, box.t - 0.2, box.w - 0.5, box.h - 0.35),
           [R("产出", color=mix(col, 0.8, "000000"), bold=True, size=12.5), P(R(text, color=INK, bold=True), before=2)],
           size=12.5, anchor="m")


STAGES = [
    (BLUE, "需求与调研", ["明确研究问题", "调研相关案例", "确定课程范围"], "需求说明"),
    (TEAL, "系统架构设计", ["划分系统层次", "明确模块职责", "梳理数据流向"], "架构方案"),
    (PURPLE, "对象与数据结构", ["设计知识对象", "定义字段与关系", "编写校验程序"], "格式说明、校验程序"),
    (ORANGE, "样例与原型实现", ["确定技术选型", "整理样例章节数据", "实现查询接口", "实现知识检索", "接入课程问答", "开发前端页面"],
     "可运行的原型系统"),
    (RED, "测试与评价", ["扩充教材知识", "检索与问答评测", "开展师生试用"], "评测结果与改进方案"),
]
STATUS = ["已完成", "已完成（上次汇报）", "本次进度", "下一阶段", None]


def body1(c):
    y0, pitch, w = 13.25, 6.03, 5.68
    c.line((1.2, y0), (29.8, y0), LINE, 5, cap=True)
    for i, ((col, title, items, out), status) in enumerate(zip(STAGES, STATUS)):
        card = Box(0.6 + i * pitch, 12.2, w, 10.0)
        x = card.l + 0.85
        c.card(card)
        c.shape("色条", MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Box(card.l, card.t, card.w, 0.32), col, radius=0.16)
        if status == "本次进度":
            c.shape("高亮边框", MSO_SHAPE.ROUNDED_RECTANGLE, card, line=col, lw=1.75, radius=0.247)
        c.line((x, y0), (x, card.t), col, 1.4)
        c.num(x, y0, str(i + 1), col, d=1.0, size=16)
        if status == "本次进度":
            c.chip("w", x + 0.75, y0, status, col, fill=col, color=WHITE)
        elif status:
            c.chip("w", x + 0.75, y0, status, ORANGE if status == "下一阶段" else GRAY)
        c.at("nw", card.l + 0.45, card.t - 0.6, f"阶段{i + 1}", size=13, color=col, bold=True)
        c.at("nw", card.l + 0.45, card.t - 1.2, title, size=18, color=INK, bold=True)
        c.text(Box(card.l + 0.45, card.t - 2.15, w - 0.8, 4.8), [P(t, bullet=col) for t in items], line=22)
        output_box(c, card, col, out)
    c.bottombar(0.35, "本次进度", "第三阶段的数据结构已写成格式说明，并编写了知识库的加载与校验程序，人工整理的数据可直接被系统读取和检查")


def body2(c):
    kb = Box(0.6, 13.8, 10.2, 10.4)
    c.card(kb)
    c.at("w", kb.l + 0.6, kb.t - 0.75, "知识库（人工依据教材整理）", size=18, color=INK, bold=True)
    c.ent(kb.cx, kb.t - 2.25, "章节", BLUE, w=7.6)
    en = c.ent(kb.cx, kb.t - 4.2, "知识实体（六类知识点）", PURPLE, w=7.6)
    fv = place("c", kb.cx, kb.t - 5.75, 5.8, 0.95)
    c.line((kb.cx, en.b), (kb.cx, fv.t), ORANGE, 1.0)
    c.ent(kb.cx, kb.t - 5.75, "公式变量（随公式保存）", ORANGE, w=5.8, h=0.95, size=13)
    c.ent(kb.cx, kb.t - 7.5, "知识关系", TEAL, w=7.6)
    c.at("c", kb.cx, kb.b + 0.65, "需要长期保存的课程知识", size=12.5, color=GRAY)

    for top, n, col, title, desc, arrow in [
        (13.8, "01", BLUE, "查询结果", "搜索结果、某个知识点的相邻知识、两个知识点之间的学习路径、局部图谱", "运行时查询"),
        (8.4, "02", PURPLE, "问答结果", "学生的问题、检索到的知识点和关系、模型的回答及引用依据", "检索后生成"),
    ]:
        box = Box(16.6, top, 13.6, 4.6)
        c.card(box)
        c.heading(box, n, title, col)
        c.text(Box(box.l + 0.6, box.t - 1.75, 12.0, 2.0), desc, line=19)
        y = box.cy + 0.2
        c.line((kb.r, y), (box.l, y), INK, 1.4, arrow=True)
        c.label((kb.r + box.l) / 2, y + 0.4, arrow)
    c.chip("c", 16.6 + 6.8, 8.4 - 4.6 + 0.55, "由后端即时生成，不写回知识库", GRAY)
    c.bottombar(0.35, "设计原则", "课程知识只保存一份，搜索、学习路径和问答都从同一份数据读取")


def body3(c):
    k1, k2 = Box(0.6, 13.8, 14.6, 11.3), Box(15.8, 13.8, 14.6, 11.3)
    c.card(k1)
    c.card(k2)
    c.heading(k1, "01", "章节不作为知识点", BLUE)
    c.text(Box(k1.l + 0.6, k1.t - 1.8, 13.2, 2.0),
           "章节只负责导航和排序，不进入图谱。每个知识点只属于一个章节，跨章节的联系用知识关系表达。", line=19)
    ents = []
    for box, chapter, name in [(Box(k1.l + 0.8, k1.t - 4.6, 5.6, 4.3), "章节：静电场", "高斯定律"),
                               (Box(k1.r - 0.8 - 5.6, k1.t - 4.6, 5.6, 4.3), "章节：时变电磁场", "麦克斯韦方程组")]:
        c.tint(box, BLUE)
        c.at("c", box.cx, box.t - 0.45, chapter, size=12.5, color=BLUE, bold=True)
        ents.append(c.ent(box.cx, box.cy - 0.35, name, PURPLE))
    c.link(ents[0], ents[1], None)
    c.label((ents[0].r + ents[1].l) / 2, ents[0].cy + 0.4, "前置知识")
    c.at("c", k1.cx, k1.b + 0.7, "共 6 个章节，暂不分层级", size=12.5, color=GRAY)

    c.heading(k2, "02", "公式单独作为知识点", ORANGE)
    c.text(Box(k2.l + 0.6, k2.t - 1.8, 13.2, 2.0),
           "公式可以被搜索、点击，单独展示变量和成立条件，并和定律、物理量、其他公式建立联系。", line=19)
    g2 = c.ent(k2.cx, k2.t - 5.3, "高斯定律", PURPLE)
    for dx, name in [(-3.4, "高斯定律积分形式"), (3.4, "高斯定律微分形式")]:
        c.link(g2, c.ent(k2.cx + dx, k2.t - 8.2, name, ORANGE), "表达为")
    c.bottombar(0.35, "对象编号", "每个对象有固定的英文编号（如 maxwell-equations），名称修改不影响引用，界面显示中文名称")


def body4(c):
    cm = Box(0.6, 13.8, 8.4, 11.3)
    c.card(cm)
    c.at("w", cm.l + 0.6, cm.t - 0.8, "共有内容", size=18, color=INK, bold=True)
    items = ["编号、名称、类型", "一句话说明", "所属章节（只有一个）", "别名（如「高斯定理」）", "标签"]
    c.text(Box(cm.l + 0.6, cm.t - 1.6, 7.4, 5.4),
           [P(t, bullet=GRAY) for t in items] + [P(R("适用条件", color=RED, bold=True), bullet=RED)],
           size=15, line=25)
    cond = place("s", cm.cx, cm.b + 0.5, 7.2, 2.3)
    c.tint(cond, RED)
    c.text(place("c", cond.cx, cond.cy, 6.6, 1.8), "公式的成立条件、方法的适用范围、定律的有效条件都记在这里，查看时位置统一",
           size=12.5, line=17, anchor="m")

    rows = [
        (BLUE, "概念", "定义、核心要点、常见误解", "电偶极子、电磁感应"),
        (PURPLE, "定律或定理", "文字表述、物理意义、常见误解", "高斯定律、唯一性定理"),
        (TEAL, "物理量", "符号、单位、标量或矢量", "电场强度（E，V/m）"),
        (ORANGE, "公式或方程", "数学表达式、变量及其含义", "法拉第定律微分形式"),
        (GREEN, "工程应用", "应用场景、关键参数", "传输线、波导"),
        (BLUE, "分析方法", "用途、主要步骤", "镜像法"),
    ]
    for i, (col, title, own, example) in enumerate(rows):
        r = Box(9.9, 13.8 - i * 1.9, 20.6, 1.6)
        c.card(r)
        x, y = r.l, r.cy
        c.num(x + 0.65, y, f"0{i + 1}", col, d=0.75, size=12)
        c.at("w", x + 1.3, y, title, size=16, color=INK, bold=True)
        c.chip("w", x + 4.7, y, "专有", col)
        c.at("w", x + 6.25, y, own)
        c.chip("w", x + 14.0, y, "例", GRAY)
        c.at("w", x + 15.0, y, example)
    c.bottombar(0.35, "两处调整", "定律与概念记录的内容相同，直接沿用概念的结构；「常见误解」单独记录，学生查询和 AI 回答都能用到")


def body5(c):
    f = Box(0.6, 13.8, 14.2, 11.3)
    c.card(f)
    c.chip("w", f.l + 0.6, f.t - 0.8, "公式", ORANGE)
    c.at("w", f.l + 2.1, f.t - 0.8, "法拉第定律微分形式", size=18, color=INK, bold=True)

    px_w, px_h = struct.unpack(">II", FORMULA_PNG.read_bytes()[16:24])
    fw, fh = px_w / 600 * 2.54, px_h / 600 * 2.54
    pic = c.shapes.add_picture(str(FORMULA_PNG), ex(f.cx - fw / 2), ey(f.t - 2.5 + fh / 2), el(fw), el(fh))
    c.n += 1
    pic.name = f"公式 {c.n}"

    c.text(Box(f.l + 0.6, f.t - 3.75, 13.0, 1.4),
           ["所属章节：时变电磁场        别名：法拉第方程", P(R("适用条件：适用于时变电磁场；介质连续", color=RED))],
           size=13, line=18)
    vt = place("s", f.cx, f.b + 0.5, 12.9, 4.6)
    c.tint(vt, ORANGE)

    rows = [("符号", "名称", "单位", "对应物理量"),
            (R("E", bold=True, latin="Times New Roman"), "电场强度", "V/m", "电场强度"),
            (R("B", bold=True, latin="Times New Roman"), "磁感应强度", "T", "磁感应强度"),
            (R("t", italic=True, latin="Times New Roman"), "时间", "s", "无（不单独建模）")]
    widths = [1.7, 3.3, 2.2, 4.9]
    gf = c.shapes.add_table(len(rows), len(widths), ex(vt.l + 0.4), ey(vt.t - 0.35), el(sum(widths)), el(3.9))
    c.n += 1
    gf.name = f"变量表 {c.n}"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.remove(tblPr.find(qn("a:tableStyleId")))
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for j, w in enumerate(widths):
        tbl.columns[j].width = el(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = el(0.9 if i == 0 else 1.0)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = 0
            cell.margin_top = cell.margin_bottom = el(0.05)
            fill_text(cell.text_frame, P(val), size=14, bold=i == 0, anchor="m")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            # 不套表格样式时部分软件会画默认网格，四条边逐一写明：只保留表头下方的横线
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ("lnL", "lnR", "lnT", "lnB"):
                fill = (f'<a:solidFill><a:srgbClr val="{TEXT}"/></a:solidFill>' if side == "lnB" and i == 0
                        else "<a:noFill/>")
                tcPr.append(parse_xml(f'<a:{side} {nsdecls("a")} w="12700">{fill}</a:{side}>'))

    for top, name, sym, unit in [(13.8, "电场强度", "E", "V/m"), (11.9, "磁感应强度", "B", "T")]:
        q = Box(16.6, top, 13.8, 1.6)
        c.card(q)
        c.chip("w", q.l + 0.4, q.cy, "物理量", TEAL)
        c.at("w", q.l + 2.2, q.cy, name, size=16, color=INK, bold=True)
        c.at("w", q.l + 7.0, q.cy, f"符号 {sym}      单位 {unit}      矢量", size=13.5)
        c.line((f.r, q.cy), (q.l, q.cy), TEAL, 1.3, arrow=True)
        c.label((f.r + q.l) / 2, q.cy + 0.38, "关联")

    for top, n, col, title, desc in [
        (9.7, "01", BLUE, "点击变量即可跳转", "学生看公式时，点击 E 或 B 就能看到对应物理量的定义和单位"),
        (5.85, "02", RED, "自动检查写法是否一致", "同一物理量出现在很多公式中，名称或单位写得不一致时自动报错"),
    ]:
        u = Box(16.6, top, 13.8, 3.55)
        c.card(u)
        c.heading(u, n, title, col, dy=0.85)
        c.text(Box(u.l + 0.6, u.t - 1.65, 12.6, 1.6), desc, line=19)
    c.bottombar(0.35, "记录原则", "公式里出现哪些物理量只通过变量关联记录，不再另建关系，同一件事只记一处")


def body6(c):
    rels = [
        (BLUE, "前置知识", "先学 → 后学", "高斯定律 → 麦克斯韦方程组", True),
        (BLUE, "定义", "定义者 → 被定义者", "电场强度定义式 → 电场强度", True),
        (ORANGE, "表达为", "定律 → 公式", "高斯定律 → 高斯定律积分形式", True),
        (PURPLE, "推导出", "已知 → 推导结果", "麦克斯韦方程组 → 波动方程", True),
        (PURPLE, "使用", "使用者 → 被使用者", "镜像法 → 唯一性定理", True),
        (TEAL, "描述", "定律或公式 → 现象", "法拉第定律 → 电磁感应", True),
        (GREEN, "应用于", "知识 → 应用", "波动方程 → 均匀平面波", True),
        (GRAY, "等价于", "无方向", "法拉第定律积分形式 — 微分形式", False),
        (GRAY, "相关", "无方向（兜底）", "电流连续性方程 — 位移电流", False),
    ]
    for i, (col, title, direction, example, directed) in enumerate(rels):
        r = Box(0.6 + (i % 3) * 10.07, 13.8 - (i // 3) * 3.75, 9.6, 3.35)
        c.card(r)
        c.chip("w", r.l + 0.45, r.t - 0.7, "有向" if directed else "无向", col)
        c.at("w", r.l + 1.95, r.t - 0.7, title, size=18, color=INK, bold=True)
        c.at("w", r.l + 0.45, r.t - 1.6, direction, color=mix(col, 0.8, "000000"), bold=True)
        c.at("w", r.l + 0.45, r.t - 2.5, f"例：{example}", size=13, color=GRAY)
    c.bottombar(0.35, "每条关系", "起点 + 终点 + 关系类型 + 一句说明（必填）；有向关系的方向由类型决定")


def body7(c):
    c.card(Box(0.6, 13.8, 29.8, 11.6))

    def pos(gx, gy):
        return 4.0 + 2.0 * gx, 7.4 + 1.45 * gy

    def chap(name, chapter):
        return [P(name), P(R(chapter, size=11.5, color=GRAY))]

    nodes = {
        "div": ((0, 3), chap("散度", "矢量分析"), BLUE),
        "gl": ((3.7, 3), chap("高斯定律", "静电场"), PURPLE),
        "mx": ((7.4, 3), chap("麦克斯韦方程组", "时变电磁场"), PURPLE),
        "we": ((11.1, 3), chap("波动方程", "电磁波"), ORANGE),
        "glf": ((3.7, 0.8), "高斯定律积分形式", ORANGE),
        "fd": ((7.4, 0.8), "法拉第定律微分形式", ORANGE),
        "upw": ((11.1, 0.8), "均匀平面波", GREEN),
        "fi": ((5.2, -1.4), "法拉第定律积分形式", ORANGE),
        "emi": ((9.7, -1.4), "电磁感应", BLUE),
        "cce": ((0, 0.8), "电流连续性方程", ORANGE),
        "dc": ((0, -1.4), "位移电流", BLUE),
    }
    boxes = {}
    for k, ((gx, gy), paras, col) in nodes.items():
        x, y = pos(gx, gy)
        name = paras if isinstance(paras, str) else paras[0].runs[0].text
        boxes[k] = place("c", x, y, max(3.4, tw(name, 14) + 0.6), 1.45 if isinstance(paras, list) else 1.15)
    edges = [("div", "gl", "前置"), ("gl", "mx", "前置"), ("mx", "we", "推导出"), ("gl", "glf", "表达为"),
             ("mx", "fd", "表达为"), ("we", "upw", "应用于"), ("fd", "emi", "描述"), ("div", "cce", "前置")]
    for a, b, t in edges:
        c.link(boxes[a], boxes[b], t)
    c.link(boxes["fi"], boxes["fd"], "等价于", arrow=False, dash=True)
    c.link(boxes["cce"], boxes["dc"], "相关", arrow=False, dash=True)
    for k, ((gx, gy), paras, col) in nodes.items():
        b = boxes[k]
        c.ent(b.cx, b.cy, paras, col, w=b.w, h=b.h)

    x0, y = 5.5, 3.25
    for i, (col, t) in enumerate([(BLUE, "概念"), (TEAL, "物理量"), (PURPLE, "定律"), (ORANGE, "公式"), (GREEN, "工程应用")]):
        x = x0 + i * 2.6
        c.shape("图例", MSO_SHAPE.ROUNDED_RECTANGLE, place("c", x, y, 0.7, 0.42), mix(col, 0.10), col, 1.0, radius=0.07)
        c.at("w", x + 0.5, y, t, size=12.5)
    c.line((x0 + 13.4, y), (x0 + 14.6, y), GRAY, 1.1, arrow=True)
    c.at("w", x0 + 14.75, y, "有方向", size=12.5)
    c.line((x0 + 17.0, y), (x0 + 18.2, y), GRAY, 1.1, dash=True)
    c.at("w", x0 + 18.35, y, "无方向", size=12.5)
    c.bottombar(0.35, "数据示例", "节点下方为所属章节，节点与关系均取自当前整理的数据，展示四个章节之间的联系")


def body8(c):
    rows = [
        (BLUE, "前置与推导分开记录", "「A 推导出 B」也意味着「A 是 B 的前置」，但系统不自动补，需要时两条都写",
         "学生问「学波动方程之前要会什么」，应得到明确答案，而不是让程序去猜"),
        (ORANGE, "无方向的关系要少用", "「等价于」只用于物理上完全等价的情况，如同一定律的积分形式与微分形式",
         "「相关」只在排除前置、推导、使用、应用之后才用，避免含义模糊的连线"),
        (TEAL, "同一事实不重复记录", "知识点属于哪一章，由所属章节说明；公式里有哪些物理量，由变量关联说明",
         "同一件事只记一处，才不会前后矛盾"),
    ]
    for i, (col, title, a, b) in enumerate(rows):
        r = Box(0.6, 13.8 - i * 3.75, 29.8, 3.3)
        c.card(r)
        c.num(r.l + 1.1, r.cy, f"0{i + 1}", col, d=1.15, size=17)
        c.at("w", r.l + 2.2, r.cy, title, size=18, w=7.2, color=INK, bold=True)
        c.line((r.l + 9.7, r.cy - 1.05), (r.l + 9.7, r.cy + 1.05), LINE, 1.0)
        c.text(place("w", r.l + 10.2, r.cy, 19.2, 2.6),
               [P(a), P(R(b, size=13, color=GRAY), before=8)], line=None, anchor="m")
    c.bottombar(0.35, "设计调整", "去掉了最初的「包含」关系：它的用法都能由所属章节、「表达为」和「相关」覆盖")


def body9(c):
    c.card(Box(0.6, 13.8, 29.8, 5.0))

    def step(x, y, label, col, w=4.4):
        box = place("c", x, y, w, 1.45)
        shp = c.shape("步骤", MSO_SHAPE.ROUNDED_RECTANGLE, box, mix(col, 0.09), col, 1.1, radius=0.21)
        fill_text(shp.text_frame, label, size=14, line=18, align="c", anchor="m", wrap=False)
        return box

    e1 = step(4.0, 11.6, ["人工整理", "课程知识"], BLUE)
    e2 = step(10.0, 11.6, ["系统启动", "读取知识库"], PURPLE)
    e3 = place("c", 15.8, 11.6, 4.2, 2.2)
    shp = c.shape("判断", MSO_SHAPE.DIAMOND, e3, mix(ORANGE, 0.09), ORANGE, 1.1)
    fill_text(shp.text_frame, "20 条检查", size=14, align="c", anchor="m", wrap=False)
    e4 = step(24.0, 12.55, "正常提供查询和问答", TEAL, w=5.4)
    e5 = step(24.0, 10.55, "报错并指出具体位置", RED, w=5.4)

    c.line(e1.east, e2.west, arrow=True)
    c.line(e2.east, (e3.l, e3.cy), arrow=True)
    bx = 19.0
    c.line((e3.r, e3.cy), (bx, e3.cy), GRAY, 1.1)
    c.poly([(bx, e3.cy), (bx, e4.cy), e4.west], TEAL)
    c.poly([(bx, e3.cy), (bx, e5.cy), e5.west], RED)
    c.label((bx + e4.l) / 2, e4.cy + 0.38, "通过")
    c.label((bx + e5.l) / 2, e5.cy + 0.38, "不通过")
    yb = 9.33
    c.poly([e5.south, (e5.cx, yb), (e1.cx, yb), e1.south], RED, dash=True)
    c.label(14.0, yb, "修改后重新启动")

    rules = [
        (BLUE, "编号是否唯一", "章节、知识点编号不重复；同一对知识点之间不重复建同一种关系"),
        (PURPLE, "引用是否存在", "所属章节、关系两端的知识点必须存在；变量只能关联物理量"),
        (ORANGE, "内容是否匹配", "专有内容与知识点类型对应；公式的数学表达式不能为空"),
        (TEAL, "取值是否一致", "变量的名称、单位与物理量一致；关系不能指向自己"),
    ]
    for i, (col, title, desc) in enumerate(rules):
        k = Box(0.6 + i * 7.55, 7.9, 7.1, 4.5)
        c.card(k)
        c.num(k.l + 0.85, k.t - 0.85, f"0{i + 1}", col)
        c.at("w", k.l + 1.6, k.t - 0.85, title, size=17, color=INK, bold=True)
        c.text(Box(k.l + 0.45, k.t - 1.7, 6.2, 2.5), desc, line=19)
    c.bottombar(0.35, "人工审核", "检查只保证格式和引用正确；内容是否符合教材、关系方向是否写反，仍需对照教材逐条确认")


def body10(c):
    c.card(Box(0.6, 13.8, 29.8, 2.4))
    for i, ((col, title, _, _), status) in enumerate(zip(STAGES, STATUS)):
        x = 2.0 + i * 6.0
        c.num(x, 12.95, str(i + 1), col, d=0.9)
        c.at("w", x + 0.65, 12.95, title, size=15, color=INK, bold=True)
        if status == "下一阶段":
            c.chip("c", x + 1.9, 12.0, status, ORANGE, size=11)
        elif status:
            c.chip("c", x + 1.9, 12.0, "已完成", GRAY, size=11)

    tech = Box(0.6, 11.1, 29.8, 1.1)
    c.card(tech)
    c.num(tech.l + 0.85, tech.cy, "01", NAVY)
    c.at("w", tech.l + 1.55, tech.cy, "确定技术选型", size=16, color=INK, bold=True)
    c.at("w", tech.l + 5.6, tech.cy, "实现前先确定后端、数据存储、检索、大模型和前端各用什么技术，并说明理由")
    c.chip("e", tech.r - 0.4, tech.cy, "产出：选型说明", NAVY)

    works = [
        (ORANGE, "整理样例数据", "选一个包含概念、定律、公式和应用的代表章节，按本次结构整理并检验", "样例章节数据"),
        (BLUE, "实现查询接口", "在样例数据上实现搜索、知识详情、相邻知识和学习路径查询", "查询接口"),
        (TEAL, "实现知识检索", "根据学生的问题找出相关的知识点和关系，作为回答的依据", "检索模块"),
        (PURPLE, "接入课程问答", "把检索结果交给大模型组织回答，并返回引用的知识点", "问答接口"),
        (GREEN, "开发前端页面", "知识搜索与详情、局部图谱展示、课程问答对话", "前端页面"),
    ]
    for i, (col, title, desc, out) in enumerate(works):
        s = Box(0.6 + i * 6.03, 9.75, 5.68, 7.55)
        c.card(s)
        c.num(s.l + 0.85, s.t - 0.85, f"0{i + 2}", col)
        c.at("w", s.l + 1.55, s.t - 0.85, title, size=16, color=INK, bold=True)
        c.text(Box(s.l + 0.45, s.t - 1.75, s.w - 0.8, 4.0), desc, line=21)
        output_box(c, s, col, out)
    c.bottombar(0.35, "存储方式", "现阶段用一个结构化文件保存，便于人工检查和版本管理；数据变大后再迁移到图数据库")


BODIES = [body1, body2, body3, body4, body5, body6, body7, body8, body9, body10]


# ============================================================ 主流程

work = Path(tempfile.mkdtemp(prefix="deck-"))
unpacked = work / "unpacked"
with zipfile.ZipFile(TEMPLATE) as z:
    z.extractall(unpacked)
content_index = build_from_template(unpacked)

staged = work / "staged.pptx"
with zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(unpacked.rglob("*")):
        if f.is_file():
            z.write(f, f.relative_to(unpacked).as_posix())

prs = Presentation(str(staged))
for idx, draw in zip(content_index, BODIES):
    draw(Canvas(prs.slides[idx]))
OUTPUT.unlink(missing_ok=True)
prs.save(str(OUTPUT))
shutil.rmtree(work)
print("written", OUTPUT)
