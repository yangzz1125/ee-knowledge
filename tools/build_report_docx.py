from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "第一次调研汇报稿.md"
OUTPUT = ROOT / "reports" / "第一次调研汇报稿.docx"

NAVY = "102A43"
BODY = "263746"
SECONDARY = "64748B"
TEAL = "18B8A6"
LINE = "D9E2EC"
HEADER_FILL = "173A5E"
ALT_FILL = "F5F8FB"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def clear_paragraph_borders(target) -> None:
    element = target._element if hasattr(target, "_element") else target._p
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, font_size: float = 10.5, color: str = BODY) -> None:
    pattern = re.compile(r"(\[[^\]]+\]\([^\)]+\)|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(SECONDARY)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.32

    title = doc.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_paragraph_borders(title)

    for name, size, before, after in (("Heading 1", 15, 18, 8), ("Heading 2", 12.5, 14, 6), ("Heading 3", 11.5, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in [s.name for s in doc.styles]:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(SECONDARY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False


def add_figure(doc: Document, alt: str, rel_path: str, caption_text: str | None) -> None:
    image_path = (ROOT / rel_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.65))
    if caption_text:
        cap = doc.add_paragraph(style="Figure Caption")
        cap.add_run(caption_text)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.45), Inches(2.35), Inches(3.05)] if len(rows[0]) == 3 else [Inches(2.1)] * len(rows[0])
    for r_idx, row_values in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if r_idx == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, ALT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, font_size=9.2 if r_idx else 9.2, color="FFFFFF" if r_idx == 0 else BODY)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build() -> None:
    doc = Document()
    style_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    paragraph_buffer: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, "".join(paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush_paragraph()
            i += 1
            continue
        if line.startswith("# "):
            flush_paragraph()
            p = doc.add_paragraph(style="Title")
            p.add_run(line[2:].replace("《", "").replace("》", ""))
            clear_paragraph_borders(p)
            i += 1
            continue
        heading = re.match(r"^(#{2,3})\s+(.*)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {level}")
            p.add_run(heading.group(2))
            i += 1
            continue
        image = re.match(r"!\[([^\]]*)\]\(([^\)]+)\)$", line)
        if image:
            flush_paragraph()
            caption_text = None
            if i + 1 < len(lines) and re.match(r"^\*.+\*$", lines[i + 1].strip()):
                caption_text = lines[i + 1].strip()[1:-1]
                i += 1
            add_figure(doc, image.group(1), image.group(2), caption_text)
            i += 1
            continue
        if line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("| ---"):
            flush_paragraph()
            table_rows: list[list[str]] = []
            table_rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].startswith("| "):
                table_rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if line.startswith("> "):
            flush_paragraph()
            p = doc.add_paragraph(style="Quote")
            p.paragraph_format.left_indent = Inches(0.25)
            add_inline(p, line[2:], font_size=9.5, color=SECONDARY)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        paragraph_buffer.append(line.strip())
        i += 1
    flush_paragraph()
    doc.core_properties.title = "电磁场与波 AI 知识图谱调研报告"
    doc.core_properties.subject = "课程知识图谱前期调研与系统架构设计"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
